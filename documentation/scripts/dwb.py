"""Document writing bridge: shared helpers for building the professional .docx."""
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLUE = RGBColor(0x1F, 0x4E, 0x79)
BLUE2 = RGBColor(0x2E, 0x74, 0xB5)
DARK = RGBColor(0x17, 0x17, 0x17)
GRAY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF2, 0xF2, 0xF2)

doc = None
FIG = 0
TAB = 0
BMARK = 0
DIAGRAM_DIR = None


def setup(diagram_dir):
    global doc, DIAGRAM_DIR
    DIAGRAM_DIR = diagram_dir
    doc = Document()
    _page_setup()
    _styles()
    return doc


def _page_setup():
    for s in doc.sections:
        s.page_width = Cm(21.0)
        s.page_height = Cm(29.7)
        s.top_margin = Cm(2.54)
        s.bottom_margin = Cm(2.54)
        s.left_margin = Cm(2.8)
        s.right_margin = Cm(2.4)
        s.footer_distance = Cm(1.27)
        s.header_distance = Cm(1.27)


def _styles():
    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(12)
    st.font.color.rgb = DARK
    pf = st.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:ascii'), 'Times New Roman')
    rfonts.set(qn('w:hAnsi'), 'Times New Roman')
    rfonts.set(qn('w:cs'), 'Times New Roman')

    def _head(name, size, color, space_before, space_after, bold=True):
        s = doc.styles[name]
        s.font.name = 'Arial'
        s.font.size = Pt(size)
        s.font.bold = bold
        s.font.color.rgb = color
        s.element.get_or_add_rPr()
        rf = s.element.rPr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts')
            s.element.rPr.append(rf)
        rf.set(qn('w:ascii'), 'Arial')
        rf.set(qn('w:hAnsi'), 'Arial')
        p = s.paragraph_format
        p.space_before = Pt(space_before)
        p.space_after = Pt(space_after)
        p.line_spacing = 1.15
        p.keep_with_next = True

    _head('Heading 1', 16, BLUE, 0, 12)
    _head('Heading 2', 14, BLUE2, 10, 6)
    _head('Heading 3', 12, BLUE, 8, 4)

    for name in ('Title', 'Subtitle'):
        try:
            s = doc.styles[name]
            s.font.name = 'Arial'
            s.font.color.rgb = BLUE
        except KeyError:
            pass


def set_update_fields():
    el = doc.settings.element
    u = OxmlElement('w:updateFields')
    u.set(qn('w:val'), 'true')
    el.append(u)


def _field(paragraph, instr, cached='', style=None, size=None, bold=False):
    run = paragraph.add_run()
    if style is not None:
        run.style = style
    if size is not None:
        run.font.size = Pt(size)
    run.font.bold = bold
    r = run._r
    b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin')
    i = OxmlElement('w:instrText'); i.set(qn('xml:space'), 'preserve'); i.text = instr
    s = OxmlElement('w:fldChar'); s.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t'); t.text = cached
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end')
    for el in (b, i, s, t, e):
        r.append(el)
    return run


def _bookmark(paragraph, name):
    sid = str(_bm_id())
    b = OxmlElement('w:bookmarkStart'); b.set(qn('w:id'), sid); b.set(qn('w:name'), name)
    e = OxmlElement('w:bookmarkEnd'); e.set(qn('w:id'), sid)
    paragraph._p.insert(0, b)
    paragraph._p.append(e)


def _bm_id():
    global BMARK
    BMARK += 1
    return BMARK


def H(text, level=1, numbered=None):
    p = doc.add_heading('', level=level)
    if level == 1:
        p.paragraph_format.page_break_before = True
    run = p.add_run(text)
    return p


def P(text='', align='justify', bold=False, italic=False, size=None, space_after=None,
      indent_first=None, color=None):
    p = doc.add_paragraph()
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    p.alignment = {'justify': WD_ALIGN_PARAGRAPH.JUSTIFY, 'left': WD_ALIGN_PARAGRAPH.LEFT,
                   'center': WD_ALIGN_PARAGRAPH.CENTER, 'right': WD_ALIGN_PARAGRAPH.RIGHT}[align]
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    if indent_first is not None:
        p.paragraph_format.first_line_indent = Cm(indent_first)
    return p


def RICH(segments, align='justify', size=None):
    """segments: list of (text, bold, italic) or (text, bold, italic, size)."""
    p = doc.add_paragraph()
    p.alignment = {'justify': WD_ALIGN_PARAGRAPH.JUSTIFY, 'left': WD_ALIGN_PARAGRAPH.LEFT,
                   'center': WD_ALIGN_PARAGRAPH.CENTER}[align]
    for seg in segments:
        txt = seg[0]
        bold = seg[1] if len(seg) > 1 else False
        italic = seg[2] if len(seg) > 2 else False
        sz = seg[3] if len(seg) > 3 else size
        r = p.add_run(txt)
        r.bold = bold
        r.italic = italic
        if sz:
            r.font.size = Pt(sz)
    return p


def B(items, style='bullet'):
    for it in items:
        if isinstance(it, tuple):
            lead, rest = it
            p = doc.add_paragraph(style='List Bullet')
            r = p.add_run(lead + ' ')
            r.bold = True
            p.add_run(rest)
        else:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(str(it))
        p.paragraph_format.space_after = Pt(2)


def NL(items):
    for i, it in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f'{i}.\t')
        r.bold = True
        p.add_run(it)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def page_break():
    doc.add_page_break()


def img(path, width_in=6.1, caption=None, bookmark=None):
    global FIG
    FIG += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(path, width=Inches(width_in))
    cap = doc.add_paragraph(style='Caption')
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(10)
    if bookmark:
        _bookmark(cap, bookmark)
    _field(cap, ' SEQ Figure \\* ARABIC ', str(FIG))
    run2 = cap.add_run(': ' + caption)
    run2.italic = True
    run2.font.size = Pt(10)
    return FIG


def cap_table(caption, bookmark=None):
    global TAB
    TAB += 1
    cap = doc.add_paragraph(style='Caption')
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(4)
    if bookmark:
        _bookmark(cap, bookmark)
    _field(cap, ' SEQ Table \\* ARABIC ', str(TAB))
    run = cap.add_run(': ' + caption)
    run.italic = True
    run.font.size = Pt(10)
    return TAB


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto')
    sh.set(qn('w:fill'), hexcolor)
    tcPr.append(sh)


def _set_cell(cell, text, size, bold=False, white=False, align='left'):
    cell.text = ''
    first = cell.paragraphs[0]
    first.paragraph_format.space_after = Pt(2)
    first.paragraph_format.space_before = Pt(2)
    first.paragraph_format.line_spacing = 1.1
    alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
                 'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
    first.alignment = alignment
    lines = str(text).split('\n')
    for i, ln in enumerate(lines):
        if i > 0:
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            p.alignment = alignment
        else:
            p = first
        r = p.add_run(ln)
        r.bold = bold
        r.italic = False
        r.font.size = Pt(size)
        if white:
            r.font.color.rgb = WHITE


def TBL(headers, rows, col_widths=None, caption=None, size=10, align=None, bookmark=None,
        first_col_left=True, zebra=True):
    """headers: list; rows: list of lists. align: list of 'left'/'center' per column."""
    if caption:
        cap_table(caption, bookmark)
    ncols = len(headers)
    tbl = doc.add_table(rows=1, cols=ncols)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    if align is None:
        align = ['left'] * ncols
    total = sum(col_widths) if col_widths else 1.0
    widths = [Inches(w) for w in col_widths] if col_widths else [Inches(6.0 / ncols)] * ncols
    hdr = tbl.rows[0].cells
    for i, h in enumerate(headers):
        _set_cell(hdr[i], h, size, bold=True, white=True, align='center')
        shade(hdr[i], '1F4E79')
        hdr[i].width = widths[i]
    for r_i, row in enumerate(rows):
        cells = tbl.add_row().cells
        for c_i, val in enumerate(row):
            _set_cell(cells[c_i], val, size, align=align[c_i] if not (c_i == 0 and first_col_left) else 'left')
            if zebra and r_i % 2 == 1:
                shade(cells[c_i], 'F2F6FB')
            cells[c_i].width = widths[c_i]
    tbl.columns[0].width = widths[0]
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def spacer(pts=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(pts)


def add_toc(title='TABLE OF CONTENTS', field='TOC \\o "1-3" \\h \\z \\u'):
    p = doc.add_paragraph()
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(14)
    r.font.name = 'Arial'
    r.font.color.rgb = BLUE
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp = doc.add_paragraph()
    _field(tp, field, 'Table of Contents will update when opened in Microsoft Word (right-click \u2192 Update Field).')
    tp.paragraph_format.line_spacing = 1.2


def add_lof(title='LIST OF FIGURES'):
    add_toc(title, 'TOC \\h \\z \\c "Figure" \\u')


def add_lot(title='LIST OF TABLES'):
    add_toc(title, 'TOC \\h \\z \\c "Table" \\u')


def add_cross_ref(paragraph, bookmark, cached, prefix='Figure ', suffix=''):
    run = paragraph.add_run(prefix)
    _field(paragraph, f' REF {bookmark} \\h ', cached)
    paragraph.add_run(suffix)


def ref_para(prefix, bookmark, cached, suffix, align='justify', prefix_bold=False):
    """Justified paragraph with a REF cross-reference field in the middle."""
    p = doc.add_paragraph()
    p.alignment = {'justify': WD_ALIGN_PARAGRAPH.JUSTIFY, 'left': WD_ALIGN_PARAGRAPH.LEFT,
                   'center': WD_ALIGN_PARAGRAPH.CENTER}[align]
    r = p.add_run(prefix)
    r.bold = prefix_bold
    _field(p, f' REF {bookmark} \\h ', str(cached))
    p.add_run(suffix)
    return p


def set_section_pgnum(section, fmt='decimal', start=None):
    sectPr = section._sectPr
    pg = sectPr.find(qn('w:pgNumType'))
    if pg is None:
        pg = OxmlElement('w:pgNumType')
        sectPr.append(pg)
    pg.set(qn('w:fmt'), fmt)
    if start is not None:
        pg.set(qn('w:start'), str(start))
    else:
        if pg.get(qn('w:start')) is not None:
            del pg.attrib[qn('w:start')]


def footer_pagenum(section, fmt='decimal', start=None, centered=True):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.text = ''
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(p, ' PAGE \\* MERGEFORMAT ', '1')
    set_section_pgnum(section, fmt, start)


def header_text(section, text, font_size=9):
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.text = ''
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(text)
    r.font.size = Pt(font_size)
    r.font.name = 'Arial'
    r.font.color.rgb = GRAY
    r.italic = True
    pb = p.paragraph_format
    pb.space_after = Pt(2)
    pPr = p._p.get_or_add_pPr()
    bd = OxmlElement('w:pBdr')
    bt = OxmlElement('w:bottom')
    bt.set(qn('w:val'), 'single')
    bt.set(qn('w:sz'), '6')
    bt.set(qn('w:space'), '1')
    bt.set(qn('w:color'), '1F4E79')
    bd.append(bt)
    pPr.append(bd)


def empty_header(section):
    header = section.header
    header.is_linked_to_previous = False
    for p in header.paragraphs:
        p.text = ''
