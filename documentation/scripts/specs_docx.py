"""Builds Screen_Specifications_and_Navigation.docx from the markdown spec files."""
import os
import sys
import re

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import dwb
from dwb import H, P, page_break, spacer, add_toc
from docx.enum.section import WD_SECTION_START

SPECS = os.path.normpath(os.path.join(HERE, '..', 'design-system', 'specs'))
SRC = {
    'Screen Specifications': 'Screen_Specifications.md',
    'Navigation Flow': 'Navigation_Flow.md',
}
OUT = os.path.join(SPECS, 'Screen_Specifications_and_Navigation.docx')


def _plain(txt):
    txt = txt.replace('**', '')
    txt = txt.replace('`', '')
    return txt


def _bullet(text):
    p = dwb.doc.add_paragraph()
    p.paragraph_format.space_after = dwb.Pt(2)
    m = re.match(r'\*\*(.+?)\*\*\s*[.\u2013:]\s*(.*)', text)
    if m and m.group(2).strip():
        r = p.add_run(m.group(1) + '. ')
        r.bold = True
        p.add_run(m.group(2))
    else:
        p.add_run(_plain(text))
    return p


def _numbered(num, text):
    p = dwb.doc.add_paragraph()
    p.paragraph_format.left_indent = dwb.Cm(1.0)
    p.paragraph_format.first_line_indent = dwb.Cm(-0.6)
    p.paragraph_format.space_after = dwb.Pt(2)
    r = p.add_run(f'{num}.\t')
    r.bold = True
    m = re.match(r'\*\*(.+?)\*\*\s*[.\u2013:]\s*(.*)', text)
    if m and m.group(2).strip():
        r2 = p.add_run(m.group(1) + '. ')
        r2.bold = True
        p.add_run(m.group(2))
    else:
        p.add_run(_plain(text))
    return p


def add_markdown(path):
    with open(path, encoding='utf-8') as fh:
        lines = fh.read().splitlines()
    in_table = None
    num = 0
    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            in_table = None
            num = 0
            continue
        if line.startswith('|'):
            if in_table is None:
                in_table = dwb.doc.add_table(rows=1, cols=len([c for c in line.strip('|').split('|')]))
                in_table.style = 'Table Grid'
                headers = [c.strip() for c in line.strip('|').split('|')]
                for i, h in enumerate(headers):
                    cell = in_table.rows[0].cells[i]
                    cell.text = ''
                    p = cell.paragraphs[0]
                    r = p.add_run(h)
                    r.bold = True
                    r.font.size = dwb.Pt(9)
                    r.font.color.rgb = dwb.WHITE
                    dwb.shade(cell, '1F4E79')
            else:
                cells = [c.strip() for c in line.strip('|').split('|')]
                if all(set(c) <= set('-: ') for c in cells):
                    continue
                row = in_table.add_row()
                for i, c in enumerate(cells):
                    if i < len(in_table.rows[0].cells):
                        cell = row.cells[i]
                        cell.text = ''
                        p = cell.paragraphs[0]
                        r = p.add_run(_plain(c))
                        r.font.size = dwb.Pt(8)
            continue
        if line.startswith('### '):
            H(_plain(line[4:]), 3)
        elif line.startswith('## '):
            H(_plain(line[3:]), 2)
        elif line.startswith('# '):
            H(_plain(line[2:]), 1)
        elif line.startswith('- '):
            _bullet(line[2:])
        elif line.startswith('> '):
            P(_plain(line[2:]), align='justify', italic=True)
        elif re.match(r'^\d+\.\s', line):
            num += 1
            _numbered(num, re.sub(r'^\d+\.\s*', '', line))
        else:
            P(_plain(line), align='justify')


def main():
    dwb.setup(SPECS)
    dwb.set_update_fields()

    sec1 = dwb.doc.sections[0]
    sec1.different_first_page_header_footer = True
    dwb.footer_pagenum(sec1, fmt='lowerRoman', start=1)

    for _ in range(4):
        P('', space_after=20)
    P('SMART CHILD HUB', align='center', bold=True, size=20, space_after=2)
    P('Screen Specifications and Navigation Flow', align='center', bold=True, size=16, space_after=10)
    P('Design System Phase 2 \u00b7 Version 1.0.0 \u00b7 August 2026', align='center', size=11, space_after=2)
    P('Companion to Design_System_Guide.docx \u00b7 22 screens \u00b7 22 routes', align='center', size=10)
    page_break()
    add_toc('TABLE OF CONTENTS')

    sec2 = dwb.doc.add_section(WD_SECTION_START.NEW_PAGE)
    dwb.footer_pagenum(sec2, fmt='decimal', start=1)
    dwb.header_text(sec2, 'Smart Child Hub \u2013 Screen Specifications and Navigation')

    for fname in SRC.values():
        add_markdown(os.path.join(SPECS, fname))
        page_break()

    dwb.doc.save(OUT)
    print('Saved document:', OUT)


if __name__ == '__main__':
    main()
