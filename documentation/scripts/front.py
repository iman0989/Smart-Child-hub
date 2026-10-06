"""Front matter: cover, declaration, approval, abstract, acknowledgements."""
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from dwb import P, RICH, B, spacer
import dwb


def _center(text, size=12, bold=False, space_after=6, color=None, italic=False):
    p = P(text, align='center', bold=bold, size=size, space_after=space_after, color=color)
    if italic:
        for r in p.runs:
            r.italic = True
    return p


def build_cover():
    _center('[UNIVERSITY NAME]', 20, True, 2)
    _center('[Faculty / School of Computing and Information Technology]', 14, False, 2)
    _center('[Department of Computer Science / Software Engineering]', 13, False, 10)
    spacer(6)
    _center('[UNIVERSITY LOGO PLACEHOLDER]', 10, False, 18, color=None, italic=True)
    spacer(6)
    _center('SOFTWARE ANALYSIS AND DESIGN (SAD)', 18, True, 2)
    _center('DOCUMENTATION', 18, True, 12)
    _center('Smart Child Daily Activity Tracking System', 24, True, 4)
    _center('An Offline Family Dashboard for Monitoring Children\u2019s Daily Habits', 13, False, 16, italic=True)
    _center('Prepared by: [Student Name]', 12, False, 0)
    _center('Registration Number: [Registration Number]', 12, False, 10)
    _center('Supervisor: [Supervisor Name]', 12, False, 0)
    _center('Programme: [BSc Software Engineering / Computer Science]', 12, False, 16)
    _center('Academic Year: 2025/2026', 12, False, 0)
    _center('August 2026', 12, True, 0)


def build_declaration():
    spacer(4)
    _center('DECLARATION', 16, True, 12)
    P('I hereby declare that this Software Analysis and Design (SAD) documentation titled '
      '\u201cSmart Child Daily Activity Tracking System\u201d is my own original work and has not been submitted '
      'for any degree or examination in any other university. All sources of information used in this '
      'document have been duly acknowledged.', align='justify')
    spacer(2)
    P('Name: ______________________________            Signature: ______________________', align='left')
    spacer(1)
    P('Date: ______________________________', align='left')
    spacer(1)
    P('Supervisor: ______________________________            Signature: ______________________', align='left')
    spacer(1)
    P('Date: ______________________________', align='left')


def build_approval():
    _center('APPROVAL PAGE', 16, True, 12)
    P('This Software Analysis and Design (SAD) documentation has been examined and is approved for '
      'submission in partial fulfilment of the requirements of the [Programme] at [University Name].',
      align='justify')
    spacer(2)
    TBL_sig = dwb.doc.add_table(rows=3, cols=2)
    TBL_sig.style = 'Table Grid'
    rows = [
        ('Supervisor', 'Name & Signature:\n\nDate:'),
        ('Head of Department', 'Name & Signature:\n\nDate:'),
        ('Dean / Coordinator', 'Name & Signature:\n\nDate:'),
    ]
    for i, (role, field) in enumerate(rows):
        c0, c1 = TBL_sig.rows[i].cells
        c0.text = role
        c1.text = field
    for row in TBL_sig.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(4)
    dwb.doc.add_paragraph()


def build_abstract():
    _center('ABSTRACT', 16, True, 12)
    P('The Smart Child Daily Activity Tracking System is an offline family dashboard designed for '
      'households in which children do not own smartphones. The system runs on a single shared tablet '
      'or touchscreen placed in the home; children simply select their own profile, while the parent '
      'administers the entire system. This document presents the complete software analysis and design '
      'of the system. The analysis began with a review of existing child habit-tracking applications, '
      'which confirmed a gap: no existing product combines shared-device operation, complete offline '
      'functionality, a parent-verification workflow, culturally aware scheduling and periodic family '
      'reporting. Requirements were elicited through a questionnaire and interviews and specified as '
      'twenty-eight functional requirements, eleven non-functional requirements and ten business '
      'rules. The design includes two actors, twelve parent use cases and nine child use cases, a '
      'context diagram, DFDs at levels zero, one and two, an entity-relationship model, a thirteen-'
      'table MySQL schema, a domain class diagram, sequence, activity and state diagrams, and five '
      'workflow diagrams covering coins, approvals, rewards and the parent and child processes. The '
      'architecture is a five-layer stack: a Flutter client, a Laravel REST API, business services, '
      'Eloquent ORM and MySQL, all communicating over the offline home network. Eight screen wireframes '
      'define the user interface. A complete project management plan schedules the work over four '
      'weeks, budgets the project at USD 280.50, and identifies risks, quality measures, '
      'communication and testing strategies. The document concludes with an implementation plan and '
      'is ready for direct development and university submission.', align='justify')
    spacer(1)
    RICH([('Keywords:  ', True), ('Offline family dashboard, child activity tracking, gamification, '
          'coin economy, parent verification, Flutter, Laravel, MySQL.', False)])


def build_acknowledgements():
    _center('ACKNOWLEDGEMENTS', 16, True, 12)
    P('I wish to express my sincere gratitude to my supervisor, [Supervisor Name], for the guidance, '
      'encouragement and constructive feedback provided throughout the preparation of this document. '
      'My appreciation also goes to my lecturers, colleagues and the parents who participated in the '
      'questionnaire and interviews that shaped the requirements of the system. Finally, I am grateful '
      'to my family for their patience, support and understanding during this project.', align='justify')
